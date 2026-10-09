import io,uuid
from datetime import timedelta
from fastapi.testclient import TestClient
from sqlalchemy import select,delete
from openpyxl import load_workbook
from pypdf import PdfReader
from app.api import app
from app.db import Session,Job,Run,RunDetail,ScenarioRecord,ScenarioVersion,now
from app.worker import process_one,recover
from app.schemas import Scenario

client=TestClient(app)
def config():return Scenario(horizon=30,service_end=30,worker_end=30,schedule=[5,20],incidents=[],replicates=2).model_dump()
def create_job():
    r=client.post('/api/experiments',json={'name':'TEST-'+str(uuid.uuid4()),'configurations':[config()]});assert r.status_code==201;return r.json()['id']
def drain(id):
    for _ in range(100):
        if client.get('/api/experiments/'+id).json()['status']=='completado':return
        process_one()
    raise AssertionError('Job did not finish')
def test_persistence_versions():
    c=config();c['name']='TEST-'+str(uuid.uuid4());r=client.post('/api/scenarios',json=c).json();id=r['id']
    c['reserve']=7;assert client.post('/api/scenarios?scenario_id='+id,json=c).json()['version']==2
    with Session() as s:
        vs=s.scalars(select(ScenarioVersion).where(ScenarioVersion.scenario_id==id).order_by(ScenarioVersion.version)).all();assert vs[0].config['reserve']==8 and vs[1].config['reserve']==7
def test_real_exports_and_detail():
    id=create_job();drain(id);j=client.get('/api/experiments/'+id).json();assert j['progress']==2
    rs=client.get('/api/experiments/'+id+'/runs').json();assert len(rs)==2
    assert client.get('/api/runs/'+rs[0]['id']+'/detail?limit=5').json()['total']>5
    csv=client.get('/api/experiments/'+id+'/export/csv');assert csv.status_code==200 and b'consumption' in csv.content
    xlsx=client.get('/api/experiments/'+id+'/export/xlsx');assert 'Réplicas' in load_workbook(io.BytesIO(xlsx.content)).sheetnames
    pdf=client.get('/api/experiments/'+id+'/export/pdf');assert len(PdfReader(io.BytesIO(pdf.content)).pages)>=2
def test_cancel_and_retry():
    id=create_job();assert client.post('/api/experiments/'+id+'/cancel').json()['status']=='cancelado'
    assert client.get('/api/experiments/'+id+'/export/csv').status_code==409
    assert client.post('/api/experiments/'+id+'/retry').json()['status']=='pendiente';drain(id)
def test_stale_recovery():
    id=create_job()
    with Session.begin() as s:
        j=s.get(Job,id);j.status='en ejecución';j.heartbeat=now()-timedelta(minutes=3)
    recover();assert client.get('/api/experiments/'+id).json()['status']=='fallido'
    assert client.post('/api/experiments/'+id+'/retry').status_code==200;drain(id)
def test_incompatible_comparison():
    c=config();d=dict(c,policy='P2',reserve=6)
    assert client.post('/api/experiments',json={'name':'Test','configurations':[c,d]}).status_code==422
def test_import_and_validation():
    from app.field import COLUMNS
    text=','.join(COLUMNS)+'\n2026-01-02,08:00,C1,suministro,2,,,\n2026-01-02,08:00,C1,interrupción,,08:10,,\n'
    req={'name':'TEST-'+str(uuid.uuid4()),'origin':'sintético','partition':'validación','text':text}
    assert client.post('/api/field/preview',json=req).json()['errors']==[]
    id=client.post('/api/field/import',json=req).json()['id']
    assert client.post('/api/field/import',json={**req,'dataset_id':id}).status_code==422
    assert client.get('/api/datasets/'+id).json()['observations'][0]['kg']==2
    j=create_job();drain(j)
    assert client.post('/api/validation',json={'dataset_id':id,'experiment_id':j}).status_code==422

def test_two_workers_claim_one_job():
    from concurrent.futures import ThreadPoolExecutor
    id=create_job()
    with ThreadPoolExecutor(max_workers=2) as pool: list(pool.map(lambda _:process_one(),range(2)))
    with Session() as s:
        rows=s.scalars(select(Run).where(Run.job_id==id)).all()
        assert len(rows)==2 and len({r.replicate for r in rows})==2

def test_observed_holdout_validation_persists():
    from app.field import COLUMNS
    prefix=','.join(COLUMNS)+'\n'
    date='2026-01-03'
    text=prefix+f'{date},08:00,C1,suministro,2,,,\n{date},08:00,C1,interrupción,,08:10,,\n{date},08:00,SISTEMA,actividad,,08:05,,\n{date},08:15,C1,recarga,,,,\n{date},08:00,SISTEMA,indicador,,,0,waste\n'
    req={'name':'TEST-'+str(uuid.uuid4()),'origin':'observado','partition':'validación','text':text}
    response=client.post('/api/field/import',json=req);assert response.status_code==200;dataset_id=response.json()['id']
    assert client.post('/api/field/import',json={**req,'name':'TEST-calibration','partition':'calibración'}).status_code==422
    id=create_job();drain(id)
    result=client.post('/api/validation',json={'experiment_id':id,'dataset_id':dataset_id}).json()
    assert next(m for m in result['metrics'] if m['metric']=='shortage')['observed']==10
    assert any(v['id']==result['id'] for v in client.get('/api/validation/history').json())

def test_cancel_running_at_replica_boundary(monkeypatch):
    import app.worker as worker
    id=create_job();original=worker.simulate;count=0
    def calculate_and_cancel(*args,**kwargs):
        nonlocal count
        count+=1;result=original(*args,**kwargs)
        if count==1:
            with Session.begin() as s: s.get(Job,id).cancel=True
        return result
    monkeypatch.setattr(worker,'simulate',calculate_and_cancel)
    assert worker.process_one()
    result=client.get('/api/experiments/'+id).json()
    assert result['status']=='cancelado' and result['progress']==1 and result['summary'] is None
    assert client.get('/api/experiments/'+id+'/export/pdf').status_code==409

def test_field_plan_persists():
    previous=client.get('/api/field/plan').json()
    try:
        assert client.post('/api/field/plan',json={'calibration_days':12,'validation_days':6}).status_code==200
        from app.db import WorkspaceSetting
        with Session() as s: assert s.get(WorkspaceSetting,'field_plan').value['calibration_days']==12
        assert client.post('/api/field/plan',json={'calibration_days':0,'validation_days':6}).status_code==422
    finally: client.post('/api/field/plan',json=previous)
